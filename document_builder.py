"""Rebuild DOCX/PDF documents from structured Hinglish blocks."""
from __future__ import annotations

import logging
import re
from pathlib import Path

from docx import Document

logger = logging.getLogger(__name__)


class DocumentBuilder:
    """Generate output documents from marker-formatted text."""

    marker_re = re.compile(r"^\[(HEADING|PARAGRAPH|LIST)\]\s*$")

    def build_docx(self, structured_text: str, output_path: str | Path) -> Path:
        """Build DOCX using marker-to-style mapping."""
        doc = Document()
        blocks = self._parse_blocks(structured_text)

        for kind, content in blocks:
            if kind == "HEADING":
                doc.add_heading(content, level=1)
            elif kind == "LIST":
                doc.add_paragraph(content, style="List Bullet")
            else:
                doc.add_paragraph(content)

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)
        logger.info("Saved DOCX: %s", output_path)
        return output_path

    def build_pdf_optional(self, structured_text: str, output_path: str | Path) -> Path:
        """Optional PDF export using reportlab if installed."""
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.pdfgen import canvas
        except Exception as exc:  # pragma: no cover
            raise RuntimeError("reportlab not installed; cannot export PDF") from exc

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        c = canvas.Canvas(str(output_path), pagesize=A4)
        y = A4[1] - 40

        for kind, content in self._parse_blocks(structured_text):
            prefix = "• " if kind == "LIST" else ""
            line = f"{prefix}{content}"
            c.drawString(40, y, line[:120])
            y -= 20
            if y <= 40:
                c.showPage()
                y = A4[1] - 40

        c.save()
        logger.info("Saved PDF: %s", output_path)
        return output_path

    def _parse_blocks(self, structured_text: str) -> list[tuple[str, str]]:
        blocks: list[tuple[str, str]] = []
        raw_blocks = [b for b in structured_text.split("\n\n") if b.strip()]
        for block in raw_blocks:
            lines = block.splitlines()
            if not lines:
                continue
            marker_match = self.marker_re.match(lines[0].strip())
            if not marker_match:
                continue
            kind = marker_match.group(1)
            content = "\n".join(lines[1:]).strip()
            if content:
                blocks.append((kind, content))
        return blocks
